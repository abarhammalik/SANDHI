import os, re, glob
import numpy as np
import pandas as pd
from openpyxl import load_workbook

root = r'data\real\semg\raw\full sEMG data in correct format'
rows = []
for folder in sorted(os.listdir(root)):
    fp = os.path.join(root, folder)
    if not os.path.isdir(fp):
        continue
    for f in sorted(os.listdir(fp)):
        p = os.path.join(fp, f)
        ext = os.path.splitext(f)[1].lower()
        m = re.search(r'(\d+)(?=\.\w+$)', f)
        sid = int(m.group(1)) if m else None
        if ext == '.csv':
            with open(p, 'r', encoding='utf-8', errors='replace') as fh:
                header = fh.readline().strip().split(',')
            with open(p, 'rb') as fb:
                nrows = sum(1 for _ in fb) - 1
        else:
            wb = load_workbook(p, read_only=True)
            ws = wb.active
            header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
            nrows = (ws.max_row - 1) if ws.max_row else None
            wb.close()
        sensors = sorted({int(x) for h in header for x in re.findall(r'EMG_IMU_(\d+)_EMG_CH', str(h))})
        rows.append(dict(folder=folder, file=f, ext=ext, subject=sid, ncols=len(header), nrows=nrows,
                         sensors=str(sensors), has_rms=any('_RMS' in str(h) for h in header),
                         has_imu=any('Acc' in str(h) for h in header)))
d = pd.DataFrame(rows)
d.to_csv(r'metadata\semg_file_inventory.csv', index=False)
print('=== FILES PER FOLDER / FORMAT ===')
print(d.groupby(['folder', 'ext']).agg(files=('file', 'size'), subj_min=('subject', 'min'), subj_max=('subject', 'max'),
      ncols=('ncols', lambda s: sorted(set(s))), rows_min=('nrows', 'min'), rows_max=('nrows', 'max')).to_string())
print('\n=== LAYOUT PATTERNS (format, sensor IDs, columns) ===')
print(d.groupby(['ext', 'sensors', 'ncols']).size().to_string())
ids = [set(g['subject']) for _, g in d.groupby('folder')]
s = pd.read_csv(r'data\real\semg\bundle\full_indicators_summary.csv', encoding='gbk')
print('\nsame subject IDs in all 3 folders:', all(x == ids[0] for x in ids), '| count:', len(ids[0]), '| range:', min(ids[0]), max(ids[0]))
print('summary "number" matches file IDs:', set(s['number']) == ids[0])

print('\n=== CLINICAL SUMMARY CHECKS ===')
print(pd.crosstab(s['group'], s['K-L grade'], margins=True).to_string())
print(pd.crosstab(s['group'], s['gender'], margins=True).to_string())
print(pd.crosstab(s['group'], s['age'], margins=True).to_string())
print(s.groupby('group')['WOMAC score'].describe().to_string())
print(s.groupby('group')['BMI'].agg(['min', 'median', 'max']).to_string())
print('missing values:', {k: int(v) for k, v in s.isna().sum().items() if v > 0})
print(s[~s['age'].isin([40, 50, 60, 70, 80])][['number', 'group', 'age', 'gender']].to_string())

print('\n=== WAVEFORM PEEK (walk task, sensor 1 CH2) ===')
w = os.path.join(root, 'walk sEMG data summary')
for pat in ('*.csv', '*.xlsx'):
    try:
        p = sorted(glob.glob(os.path.join(w, pat)))[0]
        col = 'EMG_IMU_1_EMG_CH2'
        df = pd.read_csv(p, usecols=[col]) if p.endswith('.csv') else pd.read_excel(p, usecols=[col])
        x = df[col].to_numpy(dtype=float)
        a = np.abs(x)
        print(os.path.basename(p), '| n =', len(x), '| seconds at 2000 Hz =', round(len(x) / 2000, 1))
        print('  first 12 samples:', np.round(x[:12], 5).tolist())
        print('  median|x| =', float(np.median(a)), '| p99|x| =', float(np.percentile(a, 99)), '| max|x| =', float(a.max()))
    except Exception as e:
        print(pat, 'peek failed:', type(e).__name__, e)
